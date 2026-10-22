"""OneDrive: list, search and read files, make folders and save new files (Microsoft Graph).

Never deletes, moves or overwrites anything: a new file whose name is taken gets a number added.
One-time setup (README): register an app in the Microsoft Entra admin center, put its client
ID in .env as ONEDRIVE_CLIENT_ID, then
    .venv/bin/python -m jarvis.tools.onedrive login
and enter the code it prints at microsoft.com/devicelogin (a phone is fine). The token is kept
in data/onedrive-token.json and renews itself.
"""

import io
import json
import sys
import threading
import time
from urllib.parse import quote

import requests

from jarvis import config, notify
from jarvis.tools import tool

GRAPH = "https://graph.microsoft.com/v1.0"
SCOPES = "offline_access Files.ReadWrite User.Read"
TOKEN_FILE = config.ROOT / "data" / "onedrive-token.json"
MAX_TEXT = 20000  # characters of one file passed to Claude

_lock = threading.Lock()


def _login_url(path):
    return f"https://login.microsoftonline.com/{config.ONEDRIVE_TENANT}/oauth2/v2.0/{path}"


def _save(token):
    TOKEN_FILE.parent.mkdir(exist_ok=True)
    TOKEN_FILE.touch(mode=0o600)
    TOKEN_FILE.write_text(json.dumps(token))


def _token():
    if not config.ONEDRIVE_CLIENT_ID:
        raise RuntimeError("OneDrive isn't set up yet (ONEDRIVE_CLIENT_ID in .env, then "
                           "python -m jarvis.tools.onedrive login)")
    with _lock:
        try:
            token = json.loads(TOKEN_FILE.read_text())
        except FileNotFoundError:
            raise RuntimeError("OneDrive isn't signed in yet (python -m jarvis.tools.onedrive login)")
        if token["expires_at"] - time.time() < 120:
            r = requests.post(_login_url("token"), timeout=15, data={
                "client_id": config.ONEDRIVE_CLIENT_ID, "grant_type": "refresh_token",
                "refresh_token": token["refresh_token"], "scope": SCOPES})
            if r.status_code >= 400:
                notify.post("🔑 My OneDrive sign-in has expired. At the laptop, run "
                            "`.venv/bin/python -m jarvis.tools.onedrive login` in ~/jarvis.")
                raise RuntimeError(f"OneDrive sign-in expired: {r.json().get('error_description', r.text)[:200]}")
            new = r.json()
            token.update(access_token=new["access_token"], expires_at=time.time() + new["expires_in"],
                         refresh_token=new.get("refresh_token", token["refresh_token"]))
            _save(token)
        return token["access_token"]


def _graph(method, path, raw=False, **kwargs):
    r = requests.request(method, GRAPH + path, timeout=30,
                         headers={"Authorization": f"Bearer {_token()}",
                                  **kwargs.pop("headers", {})}, **kwargs)
    if r.status_code >= 400:
        try:
            message = r.json()["error"]["message"]
        except (ValueError, KeyError):
            message = r.text[:200]
        raise RuntimeError(f"OneDrive: {message}")
    if raw:
        return r
    return r.json() if r.content else None


def _item(path):
    """Graph address of a file or folder by its path from the top of OneDrive."""
    path = path.strip().strip("/")
    return "/me/drive/root" + (f":/{quote(path)}:" if path else "")


def _entry(item, parent=""):
    kind = "folder" if "folder" in item else "file"
    entry = {"name": item["name"], "type": kind}
    if item.get("lastModifiedDateTime"):
        entry["modified"] = item["lastModifiedDateTime"][:16].replace("T", " ") + " UTC"
    if parent:
        entry["path"] = f"{parent}/{item['name']}"
    if kind == "folder":
        entry["items"] = item["folder"].get("childCount")
    else:
        entry["size_kb"] = round(item.get("size", 0) / 1024)
    return entry


@tool("List what's in a OneDrive folder.",
      path={"type": "string", "description": "Folder path from the top of OneDrive, e.g. "
            "'School/ELEC 324'; empty for the top"})
def onedrive_list(path):
    items = _graph("GET", _item(path) + "/children", params={"$top": 200})["value"]
    return [_entry(i) for i in items] or "That folder is empty."


@tool("Search the owner's OneDrive for files and folders by name or content.",
      query={"type": "string"})
def onedrive_search(query):
    found = _graph("GET", f"/me/drive/root/search(q='{quote(query.replace(chr(39), chr(39) * 2))}')",
                   params={"$top": 25})["value"]
    out = []
    for i in found:
        parent = i.get("parentReference", {}).get("path", "").split("root:", 1)[-1].strip("/")
        out.append(_entry(i, parent))
    return out or f"Nothing in OneDrive matches {query!r}."


@tool("Read the text of a file in OneDrive: Word, PDF, Jupyter notebooks, and any text or code "
      "file (Markdown, CSV, Python, MATLAB, C, LaTeX...).",
      path={"type": "string", "description": "File path from the top of OneDrive"})
def onedrive_read(path):
    data = _graph("GET", _item(path) + "/content", raw=True).content
    name = path.lower()
    if name.endswith(".docx"):
        import docx
        text = "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    elif name.endswith(".pdf"):
        import pypdf
        text = "\n".join(page.extract_text() or "" for page in pypdf.PdfReader(io.BytesIO(data)).pages)
    elif name.endswith(".ipynb"):  # a Jupyter notebook: its cells, without the outputs
        cells = json.loads(data).get("cells", [])
        text = "\n\n".join(f"[{c.get('cell_type')} cell]\n" + "".join(c.get("source", []))
                           for c in cells)
    elif b"\x00" not in data[:8192]:  # text of any kind: notes, code, CSV, LaTeX, MATLAB...
        text = data.decode(errors="replace")
    else:
        raise ValueError("Jarvis can read Word, PDF, notebooks and text or code files, not this one")
    return text[:MAX_TEXT] + ("\n[...cut off]" if len(text) > MAX_TEXT else "")


@tool("Make a folder in OneDrive (and any folders above it that don't exist yet). Fine if it "
      "already exists.",
      path={"type": "string", "description": "Folder path from the top of OneDrive, e.g. "
            "'School/ELEC 324/Lab 4'"})
def onedrive_create_folder(path):
    parts = [p for p in path.strip("/").split("/") if p.strip()]
    if not parts:
        raise ValueError("Give a folder path")
    made = []
    for n in range(len(parts)):
        here = "/".join(parts[:n + 1])
        try:
            _graph("GET", _item(here))
            continue
        except RuntimeError:
            pass
        parent = "/".join(parts[:n])
        _graph("POST", _item(parent) + "/children", json={
            "name": parts[n], "folder": {}, "@microsoft.graph.conflictBehavior": "fail"})
        made.append(here)
    if not made:
        return f"{path} already exists."
    notify.record(f"📁 Made OneDrive folder {path}")
    return f"Made {path}."


@tool("Save a new file in OneDrive: a Word document (name ending .docx, from text where lines "
      "starting '# ' or '## ' become headings) or a plain-text/Markdown file. Never replaces an "
      "existing file: if the name is taken, a number is added.",
      path={"type": "string", "description": "File path from the top of OneDrive, including the "
            "name, e.g. 'School/ELEC 324/Lab 4/Prelab.docx'"},
      text={"type": "string", "description": "The content, written out in full"})
def onedrive_save(path, text):
    if path.lower().endswith(".docx"):
        import docx
        doc = docx.Document()
        for line in text.splitlines():
            level = len(line) - len(line.lstrip("#"))
            if 0 < level <= 3 and line[level:level + 1] == " ":
                doc.add_heading(line[level + 1:], level)
            else:
                doc.add_paragraph(line)
        buffer = io.BytesIO()
        doc.save(buffer)
        data, kind = buffer.getvalue(), ("application/vnd.openxmlformats-officedocument."
                                         "wordprocessingml.document")
    else:
        data, kind = text.encode(), "text/plain"
    item = _graph("PUT", _item(path) + "/content", data=data, headers={"Content-Type": kind},
                  params={"@microsoft.graph.conflictBehavior": "rename"})
    where = item.get("webUrl", "")
    notify.record(f"📄 Saved to OneDrive: [{item['name']}]({where})" if where else
                  f"📄 Saved to OneDrive: {item['name']}")
    return f"Saved {item['name']} in OneDrive."


def login():
    """Device-code sign-in: works from any browser, including a phone."""
    if not config.ONEDRIVE_CLIENT_ID:
        raise SystemExit("Put ONEDRIVE_CLIENT_ID in .env first (see README: OneDrive setup)")
    r = requests.post(_login_url("devicecode"), timeout=15,
                      data={"client_id": config.ONEDRIVE_CLIENT_ID, "scope": SCOPES})
    r.raise_for_status()
    code = r.json()
    print(code["message"], flush=True)  # "To sign in, use a web browser to open ... enter the code ..."
    deadline = time.time() + code["expires_in"]
    while time.time() < deadline:
        time.sleep(code.get("interval", 5))
        r = requests.post(_login_url("token"), timeout=15, data={
            "client_id": config.ONEDRIVE_CLIENT_ID, "device_code": code["device_code"],
            "grant_type": "urn:ietf:params:oauth:grant-type:device_code"})
        token = r.json()
        if token.get("error") == "authorization_pending":
            continue
        if "access_token" not in token:
            raise SystemExit(f"Sign-in failed: {token.get('error_description', token)}")
        _save({"access_token": token["access_token"], "refresh_token": token["refresh_token"],
               "expires_at": time.time() + token["expires_in"]})
        me = _graph("GET", "/me")
        print(f"Signed in as {me.get('userPrincipalName')}. Saved {TOKEN_FILE}.")
        print(onedrive_list(""))
        return
    raise SystemExit("The code expired; run login again.")


if __name__ == "__main__":
    if sys.argv[1:] == ["login"]:
        login()
    else:
        print(onedrive_list(" ".join(sys.argv[1:])))
